# -*- coding: utf-8 -*-
"""👑 ÜSTAD YÖNETİM bölümü ve ÜYE GÖRÜNÜRLÜK YETKİLERİ — renkli Word kılavuzu (ekran görüntüleri gömülü)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ALTIN = RGBColor(0x8A, 0x6A, 0x12)
MOR = RGBColor(0x6A, 0x3F, 0xA0)
KOYU = RGBColor(0x1A, 0x1C, 0x20)
BEYAZ = RGBColor(0xFF, 0xFF, 0xFF)
YESIL = RGBColor(0x1B, 0x6B, 0x4A)
KIRMIZI = RGBColor(0x99, 0x2F, 0x27)
MAVI = RGBColor(0x1B, 0x4F, 0x6B)

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # ...\USTAD-SALON-KENAN-WEB
RESIM1 = os.path.join(KOK, "kod", "kilavuz-ustad-yonetim.png")
RESIM2 = os.path.join(KOK, "kod", "kilavuz-yetki-tablosu.png")
CIKTI = os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop", "USTAD-KENAN-USTAD-YONETIM-VE-YETKILER.docx")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.5)
    s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)


def serit(metin, renk="6A3FA0"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("  " + metin)
    r.font.size = Pt(14); r.bold = True; r.font.color.rgb = BEYAZ
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), renk)
    pPr.append(shd)
    return p


def tablo(basliklar, satirlar, renk="6A3FA0"):
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, b in enumerate(basliklar):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(b)
        r.bold = True; r.font.color.rgb = BEYAZ; r.font.size = Pt(10.5)
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), renk)
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for s in satirlar:
        c = t.add_row().cells
        for i, v in enumerate(s):
            c[i].text = ""
            r = c[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(10.5)
    doc.add_paragraph()
    return t


def yazi(metin, boyut=11, kalin=False, renk=None):
    p = doc.add_paragraph()
    r = p.add_run(metin.replace("**", ""))
    r.font.size = Pt(boyut); r.bold = kalin
    if renk:
        r.font.color.rgb = renk
    return p


def madde(metin, boyut=10.5):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(metin.replace("**", ""))
    r.font.size = Pt(boyut)
    return p


def resim(yol, altyazi, genislik=16.4):
    if os.path.exists(yol):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(yol, width=Cm(genislik))
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(altyazi)
        r.font.size = Pt(9.5); r.italic = True; r.font.color.rgb = KOYU
    else:
        yazi("(ekran görüntüsü bulunamadı: " + os.path.basename(yol) + ")", 9.5)


# ---------------- KAPAK ----------------
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ÜSTAD KENAN KUZUCU")
r.font.size = Pt(27); r.bold = True; r.font.color.rgb = ALTIN
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("👑 ÜSTAD YÖNETİM ve ÜYE YETKİLERİ KILAVUZU")
r.font.size = Pt(16); r.bold = True; r.font.color.rgb = MOR
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Üyeler · yüzler · tablo · “neyi görebilir / neyi göremez” yetki tablosu · kategori gizleme")
r.font.size = Pt(11); r.font.color.rgb = YESIL
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Web sitesi:  USTAD-SALON-KENAN-WEB    ·    Hazırlanma: 21 Eylül 2026")
r.font.size = Pt(10); r.italic = True

# ---------------- 1 ----------------
serit("1. BU BÖLÜM NEDİR, NEREDE DURUR?")
madde("Sol menüde, en altta **LOGO ve KÜNYE'nin hemen ALTINDA** duran **👑 ÜSTAD YÖNETİM** bölümüdür.")
madde("**Yalnızca sen girişliyken görünür.** Üyeler ve ziyaretçiler bu kategoriyi menüde hiç görmez.")
madde("Adresle açmaya çalışsalar bile \"🔒 Bu bölüm yalnızca site sahibine görünür\" uyarısı çıkar ve bölüm açılmaz.")
madde("Panele giriş: **Ctrl + Alt + Y** → PIN → Giriş. Giriş yaptığın an bu bölüm menüye eklenir ve karşına gelir.")
resim(RESIM1, "👑 ÜSTAD YÖNETİM bölümünün görünümü: sayaç kutuları, üye yüzleri ve üye tablosu")

# ---------------- 2 ----------------
serit("2. BÖLÜMDE ÜÇ KATMAN VAR")
tablo(["Katman", "Ne gösterir"],
      [["📊 Sayaç kutuları", "Toplam üye · Onaylı üye · Onay bekleyen · Yasaklı üye · Onay bekleyen yorum sayısı"],
       ["👥 YÜZLER", "Her üye, adının baş harfinden oluşan renkli bir madalyonla görünür. Altında durumu (✅ onaylı / ⏳ bekliyor / 🚫 yasaklı) ve kaç bölüm görebildiği yazar."],
       ["📋 ÜYE TABLOSU", "# · Yüz · Kullanıcı adı · E-posta · Durum · Kayıt tarihi · Son giriş · Gördüğü bölüm sayısı · İşlemler düğmeleri"]],
      "6A3FA0")
yazi("Tablodaki işlem düğmeleri: ✅ onayla · ⛔ onaylama · ⏳ beklemeye al · 🚫 yasakla · 🔑 geçici şifre üret · 🔐 yetkiler · 📧 e-posta yaz · 🗑️ sil", 10.5)

# ---------------- 3 ----------------
serit("3. BİR ÜYEYE TIKLA → “NEYİ GÖREBİLİR, NEYİ GÖREMEZ”")
yazi("Yüze ya da tablodaki satıra tıkla: o kişinin karşısına bütün bölümlerin listesi çıkar ve her birinin karşısında ✅ görebilir / 🚫 göremez yazar.", 11, True, MOR)
resim(RESIM2, "Bir üyeye tıklandığında açılan GÖRÜNÜRLÜK TABLOSU (aynı ekranda üstte üyenin kimlik kartı ve görebildiği bölüm sayacı vardır)")
tablo(["Ekranda ne var", "Açıklama"],
      [["Kimlik kartı", "Baş harf madalyonu · ad/kullanıcı adı · e-posta · durum · kayıt ve son giriş tarihi · “Görebildiği bölüm: 3 / 16” sayacı"],
       ["✅ Görebilir", "İşaretli bölümü üye görür; menüsünde çıkar, açabilir"],
       ["🚫 Göremez", "İşaretsiz bölüm üyenin menüsünde hiç çıkmaz; adresle de açamaz"],
       ["👑 işaretli satır", "Yalnızca sana görünen bölüm (ÜSTAD YÖNETİM); üyede her zaman 🚫 kalır"],
       ["İzin kutusu", "İşaretlediğin an soldaki yazı ✅/🚫 olarak değişir ve üstteki sayaç güncellenir"],
       ["💾 Yetkileri Kaydet", "Yaptığın işaretlemeleri üyeye işler; üye o an sitede geziniyorsa menüsü anında süzülür"]],
      "6A3FA0")

# ---------------- 4 ----------------
serit("4. ADIM ADIM: BİR ÜYEYE İZİN VERMEK / KISMAK", "8A6A12")
madde("Ctrl + Alt + Y → PIN → Giriş. Panel açılır.")
madde("Panelden çık (✕) ya da doğrudan sol menüden **👑 ÜSTAD YÖNETİM** bölümüne git.")
madde("Yüzler bölümünden ya da tablodan **kişiye tıkla**.")
madde("Açılan GÖRÜNÜRLÜK TABLOSU'nda istediğin bölümlerin **izin kutusunu işaretle** (görecekleri), görmeyeceklerinin işaretini kaldır.")
madde("Üstteki **Görebildiği bölüm** sayacı kaç bölüm verdiğini gösterir.")
madde("**💾 Yetkileri Kaydet** düğmesine bas. Kayıt o üyeye işler ve ekran listeye döner.")
madde("**◀️ Üye listesine dön** ile geri çıkabilirsin. Panelde de aynı işi yapabilirsin: 👥 Üyeler → satırdaki 🔐.")

# ---------------- 5 ----------------
serit("5. HAZIR SEÇENEKLER (tek tıkla işaretleme)", "8A6A12")
tablo(["Düğme", "Ne yapar"],
      [["🎵 Yalnız radyo & TV", "ÜSTAD CANLI RADYO · ÜSTAD İSLAMİ RADYO · MEKKE MEDİNE CANLI TV işaretlenir"],
       ["🕌 Yalnız İslami bölümler", "İslami radyo · Kur'an-ı Kerim · Mekke Medine TV"],
       ["🖋 Yalnız yayınlarım", "Şiirler · Makaleler · Mektuplar · Yazdığı Eserler · Okuduğum Kitaplar"],
       ["👤 Yalnız tanıtım & iletişim", "Tanıtım ve Biyografi · Eğitim ve Belgeler · Teknoloji & Dijital · Sosyal İletişim · Saç Modelleri · Logo ve Künye · İletişim ve Konum"],
       ["✅ Tümünü işaretle", "Bütün bölümler açılır (üye her şeyi görür)"],
       ["⛔ Hiçbirini işaretleme", "Hiçbir bölüm görünmez (üye boş menü görür — dikkatli kullan)"],
       ["↩️ Varsayılan (tümü açık)", "Kısıt kaldırılır; üye bütün bölümleri görür"]],
      "8A6A12")
yazi("Seçenekler yalnızca KUTULARI işaretler — kalıcı olması için mutlaka 💾 Yetkileri Kaydet'e bas.", 10.5, True, KIRMIZI)

# ---------------- 6 ----------------
serit("6. ÖRNEK: BİR ÜYENİN GÖRDÜĞÜ / GÖRMEDİĞİ", "6A3FA0")
tablo(["Bölüm", "Durum"],
      [["🕌 ÜSTAD İSLAMİ RADYO", "✅ görebilir"],
       ["📗 KUR'AN-I KERİM", "✅ görebilir"],
       ["🕋 MEKKE MEDİNE CANLI TV", "✅ görebilir"],
       ["📻 ÜSTAD CANLI RADYO", "🚫 göremez"],
       ["🕊️ ŞİİRLER / ✍️ MAKALELER / ✉️ MEKTUPLAR", "🚫 göremez"],
       ["📚 YAZDIĞI ESERLER / 📖 OKUDUĞUM KİTAPLAR", "🚫 göremez"],
       ["👤 TANITIM / 🎓 EĞİTİM / 💻 TEKNOLOJİ / 🌐 SOSYAL", "🚫 göremez"],
       ["💈 SAÇ MODELLERİ / 🖼️ LOGO ve KÜNYE", "🚫 göremez"],
       ["📞 İLETİŞİM ve KONUM", "🚫 göremez"],
       ["👑 ÜSTAD YÖNETİM", "🚫 göremez (yalnızca sana ait)"]],
      "6A3FA0")
yazi("Bu örnekte üye yalnız 3 bölüm görür. Menüsünde yalnız o üç kategori çıkar; diğerlerini adresle açmaya çalışsa \"🔒 Bu bölüm için yetkiniz yok\" uyarısını görür.", 10.5)

# ---------------- 7 ----------------
serit("7. BAŞKA BİR KATEGORİYİ DE YALNIZ SANA AÇMAK", "8A6A12")
madde("Panel → 🗂️ Kategoriler → ilgili kategorinin satırındaki **👁️ Herkese açık** düğmesine bas → **🔒 Yalnız ben** olur.")
madde("O kategori artık üyelerden ve ziyaretçilerden tamamen gizlenir (menüde çıkmaz, adresle açılmaz).")
madde("Tekrar basınca 👁️ Herkese açık olur.")
madde("Not: **Üyeye özel** kısıt koymak 👑 ÜSTAD YÖNETİM'den (kişiye tıkla), **herkesten** gizlemek 🗂️ Kategoriler'deki 🔒 ile yapılır.")

# ---------------- 8 ----------------
serit("8. GİZLİLİK VE GÜVENLİK", "992F27")
madde("Üye listesi, e-postalar ve yüzler YALNIZCA sende görünür: panelin 👥 Üyeler sekmesi ve 👑 ÜSTAD YÖNETİM bölümü.")
madde("Ziyaretçi ve üyeler hiçbir üyeyi, e-postayı göremez; onay bekleyen yorumları da göremez.")
madde("Üye listesi **veri.js dosyasına yazılmaz** (site yayınlandığında o dosyayı herkes okuyabilir). Liste tarayıcı deposunda tutulur; başka bilgisayara 💾 Yedek al (JSON) ile taşınır.")
madde("Yönetim PIN'inin fabrika değeri 1981'dir. Siteyi internete açmadan önce panelden değiştir: 💾 Yedek & Dosya → Panel şifresi (PIN).")

# ---------------- 9 ----------------
serit("9. SIK SORULANLAR")
tablo(["Soru", "Cevap"],
      [["👑 ÜSTAD YÖNETİM nerede?", "Sol menüde LOGO ve KÜNYE'nin hemen altında — yalnızca sen girişliyken görünür."],
       ["Üyeler bu bölümü görebilir mi?", "Hayır. Menüde hiç çıkmaz; adresle açılsa '🔒 yalnızca site sahibine görünür' der."],
       ["Bir üye hangi bölümleri görüyor?", "👑 ÜSTAD YÖNETİM → yüze/satıra tıkla → ✅/🚫 tablosu karşında."],
       ["Bir üyeden bölümü nasıl kapatırım?", "Kişiye tıkla → o bölümün izin kutusundaki işareti kaldır → 💾 Yetkileri Kaydet."],
       ["Üyeye tek tıkla bölüm vermek?", "🎵 / 🕌 / 🖋 / 👤 hazır seçeneklerinden birine bas → 💾 Kaydet."],
       ["Yetkiyi değiştirince üye ne olur?", "Üye o an sitede geziniyorsa menüsü anında yenilenir; yeni bölüm belirir ya da kaybolur."],
       ["Yüzler nereden geliyor?", "Üyenin adının/kullanıcı adının baş harfi, kendine özel renkte madalyon olarak çizilir."],
       ["Üyeliği nasıl onaylarım?", "Kişiye tıkla → ✅ Üyeliğini onayla; ya da tabloda ✅ düğmesi. Panelde 👥 Üyeler'de de aynı düğme var."],
       ["Şifresini unutan üye?", "Kişinin satırındaki 🔑 ile geçici şifre üretilir; ekranda bir kez gösterilir, üyeye sen iletirsin."]],
      "6A3FA0")

p = doc.add_paragraph()
r = p.add_run("\nBu kılavuz USTAD-SALON-KENAN-WEB sitesinin 👑 ÜSTAD YÖNETİM bölümü için hazırlandı.  "
              "Üstad Kenan Kuzucu · Gaziantep · 0342 329 75 44 · 0537 771 24 38")
r.font.size = Pt(9.5); r.italic = True

doc.save(CIKTI)
print("yazıldı:", CIKTI, round(os.path.getsize(CIKTI) / 1024, 1), "KB")
