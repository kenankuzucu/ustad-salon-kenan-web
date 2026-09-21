# -*- coding: utf-8 -*-
"""🎮 OYUNLAR + 📚 KİTAP EKLEME kılavuzu — renkli Word, ekran görüntüleri gömülü."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

K = r"C:\Users\kenan\OneDrive\Desktop\USTAD-SALON-KENAN-WEB"
KOD = os.path.join(K, "kod")
CIKTI = r"C:\Users\kenan\OneDrive\Desktop\USTAD-OYUNLAR-VE-KITAP-EKLEME-KILAVUZU.docx"

ALTIN = RGBColor(0xE8, 0xA3, 0x3D)
KOYU = RGBColor(0x1B, 0x1D, 0x21)
YESIL = RGBColor(0x2F, 0xA9, 0x7A)
GRI = RGBColor(0x55, 0x5A, 0x60)
MOR = RGBColor(0x8F, 0x79, 0xE8)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(1.6)
sec.left_margin = Cm(1.9); sec.right_margin = Cm(1.9)

st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def golge(hucre, renk="E8A33D", dolgu="FFF7E8"):
    tcPr = hucre._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), dolgu)
    tcPr.append(shd)


def serit(metin, renk=ALTIN, boyut=15):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(metin); r.bold = True; r.font.size = Pt(boyut); r.font.color.rgb = renk
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "18")
    bot.set(qn("w:space"), "4"); bot.set(qn("w:color"), "E8A33D")
    bdr.append(bot); pPr.append(bdr)
    return p


def yazi(metin, boyut=10.5, kalin=False, renk=None, italik=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(str(metin).replace("**", "")); r.font.size = Pt(boyut)
    r.bold = kalin; r.italic = italik
    if renk is not None: r.font.color.rgb = renk
    return p


def madde(metin, boyut=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(metin).replace("**", "")); r.font.size = Pt(boyut)
    return p


def numara(metin, boyut=10.5):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(metin).replace("**", "")); r.font.size = Pt(boyut)
    return p


def tablo(basliklar, satirlar, genislikler=None):
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, b in enumerate(basliklar):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(b); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = KOYU
        golge(hdr[i])
    for satir in satirlar:
        c = t.add_row().cells
        for i, v in enumerate(satir):
            c[i].text = ""
            r = c[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(9.5)
    if genislikler:
        for i, g in enumerate(genislikler):
            for row in t.rows: row.cells[i].width = Cm(g)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def resim(dosya, genislik=15.5, baslik=None):
    yol = dosya if os.path.isabs(dosya) else os.path.join(KOD, dosya)
    if not os.path.exists(yol):
        yazi("(görüntü bulunamadı: " + os.path.basename(yol) + ")", 9, renk=GRI, italik=True)
        return
    if baslik:
        yazi(baslik, 9.5, kalin=True, renk=GRI)
    doc.add_picture(yol, width=Cm(genislik))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ------------------------------------------------------------------ KAPAK
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("👑 ÜSTAD KENAN KUZUCU"); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = GRI
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("OYUNLAR ve KİTAP EKLEME"); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = ALTIN
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ÜSTAD SALON KENAN · Web Sitesi Kılavuzu"); r.font.size = Pt(12); r.font.color.rgb = GRI
yazi("")

resim("_kl-oyun-lobi.png", 15.5, "1) Oyun Salonu — ÜSTAD KENAN KUZUCU · OYUNLAR bölümü")

serit("0. ÖZET — BU TURDA YAPILANLAR (kısa rapor)", YESIL, 16)
tablo(["Konu", "Durum"],
      [["🎮 OYUNLAR bölümü (7 oyun, gerçek 3D)", "EKLENDİ — sol menüde SAÇ MODELLERİ'nin altında"],
       ["⚪ ÜSTADIN DAMASI (Türk daması, uçan dama, zincirli yeme, yapay zekâ)", "ÇALIŞIYOR ✔"],
       ["♞ ÜSTADIN SATRANCI (rok, terfi, geçerken alma, şah, mat, pat)", "ÇALIŞIYOR ✔"],
       ["🎲 ÜSTADIN TAVLASI (zar, kırma, kapalı nokta, bardan giriş, toplama)", "ÇALIŞIYOR ✔"],
       ["💰 ÜSTADIN MİLYONERİ (15 basamak, 50:50, seyirci, çekil)", "ÇALIŞIYOR ✔"],
       ["🌍 ÜSTADIN GENEL KÜLTÜRÜ · 🏛️ ÜSTADIN TARİH YARIŞMASI · 🧠 ÜSTADIN ZEKÂ TESTİ", "ÇALIŞIYOR ✔"],
       ["🔐 Görünürlük: yalnız sen + izin verdiğin üyeler", "KURULDU ✔ (ziyaretçi menüsünde görünmez, adresle de açılmaz)"],
       ["📚 Panelden kitap ekleme (YAZDIĞIM ESERLER)", "EKLENDİ ✔ (10 alanlı form + yardım kutusu + arama + silme)"],
       ["Menü sayacı \"… kitap\" veriden hesaplanıyor", "ÇALIŞIYOR ✔"],
       ["Sitedeki 18 bölümün tamamı açıldı, içerik dolu", "JS hatası 0 · kırık görsel 0 · yatay taşma 0 ✔"]],
      [10.0, 6.0])
yazi("Ölçülen değerler: sende menüde 18 bölüm · ziyaretçide 16 bölüm (OYUNLAR ve ÜSTAD YÖNETİM gizli) · "
     "oyun salonunda 7 kart · satrançta başlangıç 20 geçerli hamle ve bütün kural testleri doğru.", 10, renk=GRI)

serit("1. 🎮 OYUNLAR BÖLÜMÜNE NASIL GİRİLİR")
numara("Sağ üst köşedeki **🛠 YÖNETİM PANELİ** düğmesine bas (giriş yapmadıysan önce Ctrl+Alt+Y, PIN: 1981).")
numara("Sol menüden **🎮 OYUNLAR** satırına bas.")
numara("Karşına **OYUN SALONU** açılır: 7 oyun kartı. Kartın altındaki **♟ OYNA / ▶️ BAŞLA** düğmesine bas.")
yazi("Tahta oyunlarında (dama, satranç, tavla) ayrıca sağ üstteki **🎮 Lobi** düğmesiyle her an oyun listesine dönebilirsin.",
     10, renk=GRI, italik=True)

serit("2. OYUNLAR (7 OYUN — HEPSİ 3 BOYUTLU)")
tablo(["Oyun", "Tür", "Ne var içinde?"],
      [["⚪ ÜSTADIN DAMASI", "Tahta", "Gerçek ahşap tahta, torna işi pullar; uçan damalar, zincirli yeme, geri alma."],
       ["♞ ÜSTADIN SATRANCI", "Tahta", "Tam kurallı: rok, piyon terfi, geçerken alma, şah ve mat kontrolü."],
       ["🎲 ÜSTADIN TAVLASI", "Tahta", "Çekme-tavla: zar, kırma, kapalı nokta, bardan giriş, toplama."],
       ["💰 ÜSTADIN MİLYONERİ", "Yarışma", "15 soruluk ödül merdiveni, 50:50 ve seyirci jokerleri, çekilme."],
       ["🌍 ÜSTADIN GENEL KÜLTÜRÜ", "Yarışma", "10 soruluk süreli bilgi yarışması."],
       ["🏛️ ÜSTADIN TARİH YARIŞMASI", "Yarışma", "Türk ve dünya tarihi soruları."],
       ["🧠 ÜSTADIN ZEKÂ TESTİ", "Test", "Sayı dizileri, örüntüler, mantık; sonunda zekâ puanı."]],
      [4.4, 2.2, 8.9])

serit("3. TAHTA OYUNLARI NASIL OYNANIR (3D KONTROLLER)")
madde("**Döndür:** Tahtayı fareyle tutup sürükle — kamera tahtanın çevresinde döner.")
madde("**Yakınlaştır:** Fare tekerleğini kullan (mobilde parmakla sürükle).")
madde("**Taş seç:** Taşına dokun. Seçilen taş hafifçe yükselir.")
madde("**Hamle yap:** Parlayan/renkli işaretli kareye (ya da noktaya) dokun. Kırmızı işaret = taş alınacak yer.")
madde("**Zorluk:** Sağ üstteki **Zorluk** kutusundan Kolay / Orta / Zor seç.")
madde("**Geri al:** Yanlış hamle yaptıysan **↩️ Geri al** düğmesi seni bir adım geri alır.")
madde("**Yeni oyun:** **🔄 Yeni oyun** düğmesi tahtayı baştan kurar.")

resim("_kl-oyun-dama.png", 15.5, "2) ÜSTADIN DAMASI — Türk daması, ahşap tahta ve işlenmiş pullar")
resim("_kl-oyun-satranc.png", 15.5, "3) ÜSTADIN SATRANCI — tam kurallı satranç, ahşap taşlar")

serit("4. BİLGİ YARIŞMALARI (MİLYONER · KÜLTÜR · TARİH · ZEKÂ)")
madde("Soru kartı gelir; altında **A · B · C · D** dört seçenek olur.")
madde("Üstteki renkli şerit **süreyi** gösterir. Süre biterse cevap yanlış sayılır.")
madde("ÜSTADIN MİLYONERİ'nde sağda **ödül merdiveni** görünür; her doğru cevap seni bir basamak yükseltir.")
madde("**💠 50:50** iki yanlış seçeneği eler · **👥 Seyirci** oylama tahmini verir · **🚪 Çekil** o anki ödülle yarışmadan çıkarsın.")
madde("Oyun bitince puanın, doğru sayın ve **🏆 rekorun** gösterilir; **🔗 Sonucu paylaş** ile paylaşabilirsin.")
madde("ZEKÂ TESTİ sonunda tahmini bir zekâ puanı verilir (eğlence amaçlıdır).")

resim("_kl-oyun-milyoner.png", 15.5, "4) ÜSTADIN MİLYONERİ — ödül merdiveni, soru kartı ve jokerler")

serit("5. KİMLER OYUNLARI GÖRÜR? (İZİNLİ ÜYE KURALI)")
yazi("Fabrika ayarında **🎮 OYUNLAR** yalnızca **sana** görünür. Ziyaretçiler ve izinsiz üyeler menüde "
     "görmez, adresle de açamaz. İstediğin üyeye tek tek açabilirsin:", 10.5)
numara("Panel → **👥 Üyeler** sekmesi.")
numara("İzin vermek istediğin üyenin adına tıkla → **GÖRÜNÜRLÜK YETKİLERİ** tablosu açılır.")
numara("Satırlarda **🎮 OYUNLAR**'ı bul ve kutusunu işaretle.")
numara("**💾 Yetkileri Kaydet** düğmesine bas. Üye sitede yenilenir ve oyunları görür.")
yazi("Hiçbir üyeye açmak istemiyorsan dokunma; oyunlar yalnız senin kalır. "
     "Bir ara herkese açmak istersen: **🗂️ Kategoriler** → OYUNLAR satırındaki "
     "**🔐 Yalnız izinli üyeler** düğmesine bas (👁️ herkese açık hâle gelir).", 10, renk=GRI, italik=True)

serit("6. 📚 KİTAP EKLEME (YAZDIĞIM ESERLER) — ÜSTAD YÖNETİM PANELİNDE")
yazi("Yazdığın kitapları tek tek kendin ekleyebilirsin; **istediğin kadar** kitap ekleyebilirsin. Yeni bir kitap "
     "eklemek 1-2 dakika sürer.", 10.5)
numara("Sağ üstten **🛠 YÖNETİM PANELİ** düğmesine bas.")
numara("**➕ İçerik Ekle** sekmesine geç.")
numara("Üstteki **Hızlı seç** düğmelerinden **📚 YAZDIĞIM ESERLER (kitap ekle)** düğmesine bas.")
numara("Açılan **📚 YENİ KİTAP / ESER EKLE — adım adım** kutusundaki 6 adımı izle ve formu doldur.")
numara("En alttaki **💾 KAYIT ET ve YAYINA AL** düğmesine bas. Kitap anında sitede görünür.")

resim("_ekran-kitap-ekle.png", 15.5, "5) Yönetim paneli — ➕ İçerik Ekle · 📚 YAZDIĞIM ESERLER formu ve adım adım yardım kutusu")
resim("_kl-eserler.png", 15.5, "6) Sitedeki karşılığı — YAZDIĞIM ESERLER: durum etiketi, yıl · sayfa, açılır tanıtım ve önsöz")

yazi("Formdaki alanlar ve ne işe yaradıkları:", 11, kalin=True)
tablo(["Alan", "Zorunlu mu?", "Açıklama"],
      [["📕 KİTABIN ADI", "Evet", "Kitabın adı. Yazıp kaydetmen yeterli."],
       ["Kitabın durumu", "Hayır", "✅ Yayında · 🖨️ Baskıda · 📝 Hazırlanıyor · 🖋️ Yazılıyor. Kartta renkli etiket olarak görünür."],
       ["Türü", "Hayır", "Roman · İnceleme · Şiir · Senaryo · Deneme…"],
       ["Yazım / yayın yılı", "Hayır", "Örn. 2026. Kartta 🗓️ olarak görünür."],
       ["Sayfa sayısı", "Hayır", "Örn. 184 → kartta 📄 184 sayfa yazar."],
       ["Kapak görseli", "Hayır", "Kapağın resmini sitenin foto klasörüne koy, adını yaz: foto/kitap-adi.jpg. Boş bırakırsan şık bir kapak çizilir."],
       ["KISA TANITIM", "Hayır", "1-2 cümle; kartın üstünde görünür."],
       ["UZUN TANITIM", "Hayır", "Arka kapak yazısı. Paragrafları boş satırla ayır; kartta '📖 Detaylı tanıtım' açılır kutusuna girer."],
       ["ÖNSÖZ", "Hayır", "Kitabın önsözü; karttaki 📝 ÖNSÖZ kutusunda açılır."],
       ["Satın alma / okuma bağlantısı", "Hayır", "Varsa adres; kartta '🔗 Kitaba git' düğmesi olur."]],
      [4.2, 2.0, 9.3])

serit("7. EKLENEN KİTABI DÜZENLEME / SİLME / ARAMA")
madde("Panelde aynı ekranda **Mevcut kayıtlar: 📚 YAZDIĞIM ESERLER (… adet)** listesi durur.")
madde("Yanlış girdiysen kitabı bul, **🗑️** düğmesiyle sil ve yeniden ekle.")
madde("Kalabalıklaşınca sağdaki **🔍 Listede ara** kutusuna kitabın adını yaz; liste anında süzülür.")
madde("Her kaydın yanında **panelden · Hazırlanıyor · 2026** gibi etiketler görünür — böylece hangi kitabı ne zaman eklediğini görürsün.")
madde("Kayıtlar bu bilgisayarın tarayıcı deposunda tutulur. **💾 Yedek & Dosya** sekmesinden yedek alırsan kaybolmaz.")

serit("8. SIK SORULANLAR")
tablo(["Soru", "Cevap"],
      [["Kaç kitap ekleyebilirim?", "Sınır yok. Her eklediğin kitap anında sitede görünür; sol menüdeki '… kitap' sayacı kendiliğinden artar."],
       ["Kapak resmi nasıl olmalı?", "Sitenin foto klasörüne koy (örn. foto/ahlaksiz-toplumlar.jpg) ve adını forma yaz. Kapak koymazsan üzerinde roman numarası olan şık bir kapak çizilir."],
       ["Oyunlarda yapay zekâ zor mu?", "Zorluk kutusundan Kolay / Orta / Zor seçebilirsin. Kolay seviyede bazen hata yapar, Zor seviyede iyi oynar."],
       ["Oyunlar kimlere görünür?", "Fabrika ayarında yalnız sana. İstediğin üyeye 👥 Üyeler → GÖRÜNÜRLÜK YETKİLERİ → OYUNLAR işaretini koyarak açabilirsin."],
       ["Rekorum kaydediliyor mu?", "Evet; her oyunun rekoru bu cihazda saklanır ve oyun salonundaki kartta görünür."],
       ["Mobil telefonla oynanır mı?", "Evet. Tahtayı parmakla sürükleyip döndürebilir, taşları dokunarak oynayabilirsin."]],
      [4.8, 10.7])

yazi("")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("👑 ÜSTAD SALON KENAN · Kenan Kuzucu · Gaziantep / Şehitkamil")
r.font.size = Pt(10); r.font.color.rgb = GRI

doc.save(CIKTI)
print("yazıldı:", CIKTI, round(os.path.getsize(CIKTI) / 1024, 1), "KB")
