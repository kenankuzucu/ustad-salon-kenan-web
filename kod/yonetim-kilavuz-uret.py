# -*- coding: utf-8 -*-
"""ÜSTAD KENAN KUZUCU WEB SİTESİ — YÖNETİM PANELİ ve ÜYELİK KILAVUZU (renkli Word)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ALTIN = RGBColor(0x8A, 0x6A, 0x12)
KOYU = RGBColor(0x1A, 0x1C, 0x20)
BEYAZ = RGBColor(0xFF, 0xFF, 0xFF)
YESIL = RGBColor(0x1B, 0x6B, 0x4A)
KIRMIZI = RGBColor(0x99, 0x2F, 0x27)
MAVI = RGBColor(0x1B, 0x4F, 0x6B)

MASaustu = os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop")
CIKTI = os.path.join(MASaustu, "USTAD-KENAN-YONETIM-VE-UYELIK-KILAVUZU.docx")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(1.7); s.bottom_margin = Cm(1.5)
    s.left_margin = Cm(1.9); s.right_margin = Cm(1.9)


def serit(metin, renk="1F6F5C"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("  " + metin)
    r.font.size = Pt(14); r.bold = True; r.font.color.rgb = BEYAZ
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), renk)
    pPr.append(shd)
    return p


def tablo(basliklar, satirlar, renk="1F6F5C"):
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, b in enumerate(basliklar):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(b)
        r.bold = True; r.font.color.rgb = BEYAZ; r.font.size = Pt(10.5)
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), renk)
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for s in satirlar:
        c = t.add_row().cells
        for i, v in enumerate(s):
            c[i].text = ""
            r = c[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(10.5)
    doc.add_paragraph()
    return t


def yazi(metin, boyut=11, kalin=False, renk=None, girinti=0.0):
    p = doc.add_paragraph()
    if girinti:
        p.paragraph_format.left_indent = Cm(girinti)
    r = p.add_run(metin)
    r.font.size = Pt(boyut); r.bold = kalin
    if renk:
        r.font.color.rgb = renk
    return p


def madde(metin, boyut=11):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(metin)
    r.font.size = Pt(boyut)
    return p


# ---------------- KAPAK ----------------
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ÜSTAD KENAN KUZUCU")
r.font.size = Pt(28); r.bold = True; r.font.color.rgb = ALTIN
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("YÖNETİM PANELİ ve ÜYELİK KILAVUZU")
r.font.size = Pt(16); r.bold = True; r.font.color.rgb = KOYU
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Panel girişi · PIN · üyeler ve yönetim onayı · üst çubuk düğmeleri · kısayollar · gizlilik")
r.font.size = Pt(11); r.font.color.rgb = YESIL
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Web sitesi:  USTAD-SALON-KENAN-WEB    ·    Hazırlanma: 21 Eylül 2026")
r.font.size = Pt(10); r.italic = True

# ---------------- 1 ----------------
serit("1. YÖNETİM PANELİNE NASIL GİRERİM? (ÜÇ YOL)")
yazi("Ziyaretçiler bu perdeyi göremez — üst çubukta hiçbir “yetkili girişi” düğmesi yoktur. Sen şu üç yoldan birini kullan:", 11, False, YESIL)
madde("1) EN KOLAY YOL — KLAVYEDEN:  Ctrl + Alt + Y  tuşlarına birlikte bas.  Karşına “🔐 Yönetim Girişi” perdesi gelir.")
madde("PIN alanına  1981  yaz ve “Giriş” düğmesine bas. Panel açılır.")
madde("Panel açıldıktan sonra üst şeritte “🛠 YÖNETİM PANELİ” düğmesi kendiliğinden belirir; bir daha tek tıkla oradan açabilirsin.")
madde("“Bu bilgisayarda beni hatırla” kutusunu işaretlersen bu bilgisayarda bir daha PIN sormaz.")
madde("2) ADRESTEN:  tarayıcı adresinin sonuna  #yonetim  yaz  →  ...\\index.html#yonetim  →  aynı PIN perdesi açılır.  (#panel ve #sahip de çalışır.)")
madde("3) İÇERİDE:  Panel açıkken sekmelerden gezerek yönetirsin (aşağıda hangi sekme ne işe yarıyor yazıyor).")
yazi("PIN’i değiştirmek için:  Panel → “💾 Yedek & Dosya” sekmesi → “Panel şifresi (PIN) — şu an: 1981” satırı → yeni PIN yaz (4-8 hane) → “PIN’i Değiştir”.", 11, True, KIRMIZI)

# ---------------- 2 ----------------
serit("2. ÜST ÇUBUKTA NE VAR, KİM NE GÖRÜR", "1B4F6B")
tablo(["Kim bakıyor", "Üst araç şeridinde gördüğü düğmeler"],
      [["Girişsiz ziyaretçi", "Tema noktaları · 🔤 Yazı Tipi · 🌌 3D sahne · 🤝 ÜYE OL · 🔗 Bağlantı"],
       ["Üye girişi yapmış kişi", "🤝 ÜYE OL kaybolur, yerine  👤 kullanıcı adı  gelir (üyelik kartını açar)"],
       ["Site sahibi (sen) girişliyken", "🤝 ÜYE OL gizlenir, yerine  🛠 YÖNETİM PANELİ  düğmesi eklenir"]],
      "1B4F6B")
yazi("Önemli: “MİSAFİR” ve “🔐 YETKİLİ GİRİŞİ” düğmeleri kaldırıldı. Sen giriş yapmadığın sürece “YÖNETİM PANELİ” yazısı sayfaya hiç eklenmez — ziyaretçi kaynak kodunda bile göremez.", 10.5, False, KIRMIZI)

# ---------------- 3 ----------------
serit("3. KİM NE YAPABİLİR (YETKİ TABLOSU)", "8A6A12")
tablo(["İşlem", "Ziyaretçi", "Onaylı üye", "Site sahibi (sen)"],
      [["Kategorileri, yazıları okuma", "VAR", "VAR", "VAR"],
       ["Yorum yazma", "YOK", "VAR (senin onayınla yayımlanır)", "VAR"],
       ["Kategoriye yazı/şiir/makale ekleme", "YOK", "YOK", "VAR"],
       ["Kategori açma, gizleme, renk değiştirme", "YOK", "YOK", "VAR"],
       ["Üye listesini görme", "YOK", "YOK", "VAR (👥 Üyeler sekmesi)"],
       ["3D sahne, tema, panel rengi seçme", "Kendi tarayıcısında", "Kendi tarayıcısında", "VAR (🎨 Görünüm & 3D)"]],
      "8A6A12")

# ---------------- 4 ----------------
serit("4. ÜYELİK: KAYIT NASIL OLUYOR, ONAY NASIL VERİLİYOR?")
yazi("Kayıt alanları: kullanıcı adı · e-posta (Gmail) · şifre (+ şifre tekrar) · istersen ad soyad.  TELEFON NUMARASI İSTENMEZ.", 11, True, YESIL)
madde("Kişi sağ üstteki  🤝 ÜYE OL  düğmesine basar, formu doldurup “✅ KAYDIMI GÖNDER” der.")
madde("Kayıt  “⏳ yönetici onayı bekliyor”  durumuna düşer. Bu kişi giriş yapmayı denerse “hesabın yönetici onayı bekliyor” mesajını görür; giriş yapamaz.")
madde("Sen panelde  👥 Üyeler  sekmesine girersin; onay bekleyen kişi en üstte görünür.")
madde("Yeşil ✅ düğmesine basarsın → kişi  onaylı üye  olur ve kullanıcı adı (veya e-postası) + şifresi ile giriş yapıp yorum yazabilir.")
madde("“⛔” onaylamama, “⏳” beklemeye alma, “🚫” yasaklama, “🔑” geçici şifre üretme, “📧” e-posta yazma, “🗑️” silme demektir.")
yazi("Şifre güvenliği: şifreler düz metin olarak saklanmaz. Her üyeye rastgele bir tuz verilir ve şifre 5000 turlu SHA-256 ile özete çevrilir. 5 hatalı girişte giriş 60 saniye kilitlenir.", 10.5)

# ---------------- 5 ----------------
serit("5. YÖNETİM PANELİ SEKMELERİ (8 SEKME)", "1B4F6B")
tablo(["Sekme", "Ne işe yarar"],
      [["➕ İçerik Ekle", "Şiir, makale, mektup, eser, kitap, sosyal iletişim, belge, saç modeli ekleme"],
       ["🗂️ Kategoriler", "Kategori açma, adını/simgesini/rengini değiştirme, sırasını oynatma, gizleme, geri getirme"],
       ["🎨 Görünüm & 3D", "3D sahne seçimi (28 sahne) · rastgele sahne · tema (30) · yazı tipi (16) · her kategorinin rengi · panelin kendi rengi · kısayol listesi"],
       ["💬 Yorumlar", "Yorumları onaylama (yayına alma), gizleme, silme"],
       ["👥 Üyeler", "Üye onayı, red, yasaklama, geçici şifre, arama, e-posta listesini kopyalama, Excel/Word indirme, üyelik ayarları"],
       ["🗂️ Kayıt Defteri", "Salon müşteri/kayıt defteri (elle kişi ekleme, durum takibi, Excel/Word)"],
       ["📣 Yayın Kaydı", "Nerede ne yayımlandı takibi"],
       ["💾 Yedek & Dosya", "Yedek al/yükle, veri.js dosyasına yazma, PIN değiştirme"]],
      "1B4F6B")

# ---------------- 6 ----------------
serit("6. KISAYOLLAR (KLAVYE)", "8A6A12")
tablo(["Kısayol", "İş"],
      [["Ctrl + Alt + Y", "Yönetim panelini aç / kapat (gizli sahip girişi — ziyaretçi bu düğmeyi görmez)"],
       ["Ctrl + Alt + G", "Panelde Görünüm ve 3D sekmesini aç"],
       ["Ctrl + Alt + U", "Üyelik penceresi (kayıt · giriş)"],
       ["Ctrl + Alt + K", "Kategoriler listesini aç"],
       ["Ctrl + Alt + S", "3D sahne panelini aç / kapat"],
       ["Ctrl + Alt + R", "Rastgele 3D sahne"],
       ["Ctrl + Alt + T", "Sıradaki tema"],
       ["Ctrl + Alt + F", "Sıradaki yazı tipi"],
       ["Ctrl + Alt + B", "Radyo bölümüne git"],
       ["Ctrl + Alt + A", "Arama kutusuna git"],
       ["Esc", "Açık pencereyi kapat"]],
      "8A6A12")

# ---------------- 7 ----------------
serit("7. GİZLİLİK VE GÜVENLİK (önemli)", "992F27")
madde("Üye listesi, e-postalar ve kimin ne yazdığı YALNIZCA sende görünür: panelin 👥 Üyeler sekmesi ve sol menüdeki 👑 ÜSTAD YÖNETİM bölümü. Ziyaretçi ve üyeler hiçbir üyeyi göremez.", 10.5)
madde("Üyelik kartındaki sayılar (kaç onaylı / kaç bekleyen) yalnızca sen girişliyken görünür; ziyaretçiye gösterilmez.", 10.5)
madde("Onay bekleyen yorumlar yalnızca sana görünür; ziyaretçi sadece yayımlanmış yorumları okur.", 10.5)
madde("Üye listesi veri.js dosyasına YAZILMAZ. (O dosya siteyi yayına verdiğinde herkes okuyabilir.) Liste tarayıcı deposunda tutulur; başka bilgisayara “💾 Yedek al (JSON)” ile taşınır.", 10.5)
madde("Yönetim PIN’inin fabrika değeri 1981’dir ve kodun içinde yazılıdır. Siteyi internete açmadan önce panelden PIN’i değiştir.", 10.5)
madde("Site şu an bilgisayarındaki dosya olarak çalışıyor; kayıtlar bu bilgisayarın tarayıcı deposunda duruyor. Başka insanların internetten üye olması ve kayıtların sana düşmesi için cPanel tarafında küçük bir sunucu (PHP) gerekir.", 10.5)

# ---------------- 8 ----------------
serit("8. 👑 ÜSTAD YÖNETİM BÖLÜMÜ (YALNIZ SANA GÖRÜNEN)", "6A3FA0")
yazi("Sol menüde LOGO ve KÜNYE'nin HEMEN ALTINDA duran 👑 ÜSTAD YÖNETİM bölümü yalnızca sen girişliyken görünür.", 11, True, ALTIN)
madde("Üyeler ve ziyaretçiler bu kategoriyi menüde görmez; adresle açmak isteseler bile '🔒 Bu bölüm yalnızca site sahibine görünür' uyarısı çıkar.", 10.5)
madde("Bölüme girdiğinde en üstte sayaç kutuları: Toplam üye · Onaylı üye · Onay bekleyen · Yasaklı · Onay bekleyen yorum.", 10.5)
madde("Altında 👥 YÜZLER: her üye, adının baş harfinden oluşan renkli madalyonla görünür (durumu ve kaç bölüm görebildiği altında yazar).", 10.5)
madde("Sonra 📋 ÜYE TABLOSU: # · Yüz · Kullanıcı adı · E-posta · Durum · Kayıt · Son giriş · Gördüğü bölüm · İşlemler (✅ ⛔ ⏳ 🚫 🔑 🔐 📧 🗑️).", 10.5)
madde("Yüze veya tablodaki satıra tıkla → karşına GÖRÜNÜRLÜK TABLOSU çıkar (aşağıda).", 10.5)

# ---------------- 9 ----------------
serit("9. BİR ÜYEYE TIKLAYINCA: NEYİ GÖREBİLİR / NEYİ GÖREMEZ", "6A3FA0")
yazi("Kişiye tıkladığın an açılan tabloda her bölümün karşısında ✅ görebilir ya da 🚫 göremez yazar.", 11, True, ALTIN)
madde("İzin sütunundaki kutuyu işaretlersen o bölüm o üyeye görünür; boş bırakırsan üye o bölümü menüde hiç görmez ve adresle de açamaz.", 10.5)
madde("Kutuyu işaretlediğin an soldaki yazı ✅/🚫 olarak değişir ve üstteki 'Görebildiği bölüm' sayacı güncellenir.", 10.5)
madde("Hazır seçenekler: 🎵 Yalnız radyo & TV · 🕌 Yalnız İslami bölümler · 🖋 Yalnız yayınlarım · 👤 Yalnız tanıtım & iletişim · ✅ Tümünü işaretle · ⛔ Hiçbirini işaretleme · ↩️ Varsayılan (tümü açık).", 10.5)
madde("💾 Yetkileri Kaydet'e basınca kayıt üyeye işler; üye o an sitede geziniyorsa menüsü anında yenilenir.", 10.5)
madde("👑 ÜSTAD YÖNETİM satırı her zaman 🚫 göremez kalır — o bölüm yalnızca sana aittir.", 10.5)
yazi("İpucu: Paneldeki 👥 Üyeler sekmesinde de her üyenin satırındaki 🔐 düğmesi aynı yetki ekranını açar; iki yerden de aynı sonucu alırsın.", 10.5)

# ---------------- 10 ----------------
serit("10. BAŞKA BİR KATEGORİYİ DE YALNIZ SANA AÇMAK", "6A3FA0")
madde("Panel → 🗂️ Kategoriler → ilgili kategorinin satırındaki 👁️ Herkese açık düğmesine bas: 🔒 Yalnız ben olur ve o kategori üyelerle ziyaretçilerden gizlenir. Tekrar basınca herkese açılır.", 10.5)
madde("Not: Üyeler için ayrı ayrı kısıt koymak 👑 ÜSTAD YÖNETİM sayfasından, kategoriyi tamamen gizlemek 🗂️ Kategoriler'deki 🔒 ile yapılır.", 10.5)

# ---------------- 11 ----------------
serit("11. SIK SORULANLAR")
tablo(["Soru", "Cevap"],
      [["PANELE NEREDEN GİRECEM?", "Ctrl + Alt + Y  →  PIN 1981  →  Giriş.  (Kısayolu unutursan adresin sonuna #yonetim yaz.)"],
       ["Şifreyi (PIN) nasıl değiştiririm?", "Panel → 💾 Yedek & Dosya → “Panel şifresi (PIN)” satırı → yeni PIN → PIN’i Değiştir."],
       ["Üye kaydını nereden onaylarım?", "Panel → 👥 Üyeler → onay bekleyen kişinin karşısındaki yeşil ✅ düğmesi."],
       ["Üye şifresini unuttu, ne yapacağım?", "Kişi “🔑 Şifremi unuttum” der; istek panele düşer. Sen 👥 Üyeler’de 🔑 düğmesine basıp geçici şifre üretir, WhatsApp/e-posta ile gönderirsin."],
       ["Yeni kaydı kapatabilir miyim?", "Evet. Panel → 👥 Üyeler → “✅ Yeni kayıt açık” düğmesi (kapatır/açar)."],
       ["Yorumlar onayımı beklesin mi?", "Panel → 👥 Üyeler → “✅ Yorum onaydan sonra yayımlanır” düğmesi ile değiştirilir."],
       ["3D sahneyi nereden değiştiririm?", "Panel → 🎨 Görünüm & 3D → 28 sahneden birine bas; ya da Ctrl + Alt + R ile rastgele."],
       ["Ziyaretçi üye listesini görebilir mi?", "Hayır. Liste yalnızca senin panelinde ve 👑 ÜSTAD YÖNETİM bölümünde."],
       ["ÜSTAD YÖNETİM bölümü nerede?", "Sol menüde, LOGO ve KÜNYE'nin hemen altında — yalnızca sen girişliyken görünür."],
       ["Bir üye hangi bölümleri görüyor, nasıl bakarım?", "👑 ÜSTAD YÖNETİM → yüze/satıra tıkla → '🔐 NEYİ GÖREBİLİR, NEYİ GÖREMEZ' tablosu açılır."],
       ["Bir üyeden bölümü nasıl kapatırım?", "👑 ÜSTAD YÖNETİM → kişiye tıkla → ilgili bölümün kutusundaki işareti kaldır → 💾 Yetkileri Kaydet."],
       ["Bütün kategorileri sadece ben görebilir miyim?", "Evet: 🗂️ Kategoriler → satırdaki 👁️ düğmesi (🔒 Yalnız ben olur)."]],
      "1F6F5C")

p = doc.add_paragraph()
r = p.add_run("\nBu kılavuz USTAD-SALON-KENAN-WEB sitesi için hazırlandı.  Üstad Kenan Kuzucu · Gaziantep · 0342 329 75 44 · 0537 771 24 38")
r.font.size = Pt(9.5); r.italic = True

doc.save(CIKTI)
print("yazıldı:", CIKTI, round(os.path.getsize(CIKTI) / 1024, 1), "KB")
