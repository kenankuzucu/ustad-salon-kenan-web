# -*- coding: utf-8 -*-
"""BİYOGRAFİ WORD BELGESİ — site verisinden üretir (elle kopyalama yok).
Kullanım:  node kod/veri-disa.js veri.js > kod/veri.json && python kod/biyografi-docx.py
Çıktı:     Masaüstü klasöründe USTAD-KENAN-KUZUCU-BIYOGRAFI.docx
"""
import json
import os
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
VERI_JSON = os.path.join(KOK, "kod", "veri.json")
CIKTI = r"C:\Users\kenan\OneDrive\Desktop\USTAD-KENAN-KUZUCU-BIYOGRAFI.docx"

ALTIN = RGBColor(0x8A, 0x6A, 0x12)
KOYU = RGBColor(0x1A, 0x1C, 0x20)
BEYAZ = RGBColor(0xFF, 0xFF, 0xFF)
YESIL = RGBColor(0x1B, 0x6B, 0x4A)

# veri.json yoksa üret
if not os.path.exists(VERI_JSON):
    src = open(os.path.join(KOK, "veri.js"), encoding="utf-8").read()
    node = r"C:\Users\kenan\AppData\Local\hermes\node\node.exe"
    r = subprocess.run([node, os.path.join(KOK, "kod", "veri-disa.js"), os.path.join(KOK, "veri.js")],
                       capture_output=True, text=True, encoding="utf-8")
    open(VERI_JSON, "w", encoding="utf-8").write(r.stdout)

D = json.load(open(VERI_JSON, encoding="utf-8"))
kisi = D["KISI"]

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(1.7); s.bottom_margin = Cm(1.5)
    s.left_margin = Cm(2); s.right_margin = Cm(2)


def serit(metin, renk="1F6F5C", boy=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("  " + metin)
    r.font.size = Pt(boy); r.bold = True; r.font.color.rgb = BEYAZ
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), renk)
    p._p.get_or_add_pPr().append(shd)
    return p


def satir_yaz(metin, kalin=False, renk=None, boy=11, aralik=4, hiza=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(aralik)
    if hiza:
        p.alignment = hiza
    r = p.add_run(metin)
    r.bold = kalin; r.font.size = Pt(boy)
    if renk:
        r.font.color.rgb = renk
    return p


def tablo(basliklar, satirlar, renk="1F6F5C"):
    t = doc.add_table(rows=1, cols=len(basliklar)); t.style = "Table Grid"
    for i, b in enumerate(basliklar):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(b); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = BEYAZ
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), renk)
        c._tc.get_or_add_tcPr().append(shd)
    for s in satirlar:
        c = t.add_row().cells
        for i, v in enumerate(s):
            c[i].text = ""
            r = c[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(10.5)
    doc.add_paragraph()
    return t


# ---------------- KAPAK ----------------
satir_yaz("ÜSTAD KENAN KUZUCU", True, ALTIN, 28, 2, WD_ALIGN_PARAGRAPH.CENTER)
satir_yaz(" • ".join(D["BIO_BASLIK"]), False, YESIL, 11, 10, WD_ALIGN_PARAGRAPH.CENTER)
# portre fotoğrafı (foto klasöründen)
PORTRE = os.path.join(KOK, "foto", "portre-resmi.jpg")
if os.path.exists(PORTRE):
    _pp = doc.add_paragraph(); _pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _pp.add_run().add_picture(PORTRE, width=Cm(5.2))
    doc.add_paragraph()

satir_yaz("RESMÎ BİYOGRAFİ", True, KOYU, 14, 4, WD_ALIGN_PARAGRAPH.CENTER)
satir_yaz("Kenan Kuzucu'nun kendi kaleme aldığı metin, olduğu gibi.", False, KOYU, 10.5, 14, WD_ALIGN_PARAGRAPH.CENTER)

# ---------------- BÖLÜMLER ----------------
for b in D["BIO_TAM"]:
    serit(b["bas"]) if b["bas"] != "YAŞAM FELSEFESİ" else serit(b["bas"], "8A6A12")
    for par in b.get("paragraf", []):
        satir_yaz(par, False, KOYU)
    if b.get("cift"):
        tablo(["Bilgi", "Değer"], [[c[0], c[1]] for c in b["cift"]])
    if b.get("madde"):
        for m in b["madde"]:
            p = doc.add_paragraph(m, style="List Bullet")
            p.runs[0].font.size = Pt(10.5)
    if b.get("alinti"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("“" + b["alinti"] + "”")
        r.italic = True; r.bold = True; r.font.size = Pt(14); r.font.color.rgb = ALTIN
    for par in b.get("devam", []):
        satir_yaz(par, False, KOYU)

satir_yaz(D.get("BIO_DIPNOT", ""), False, YESIL, 9.5, 8)

# ---------------- EK: SALON ----------------
serit("EK — SALON, HİZMETLER VE İLETİŞİM", "4B2E83")
satir_yaz(kisi["isletme"] + " — " + D["ILETISIM"]["adres"], True, KOYU, 11)
tablo(["Hizmet", "Açıklama", "Ücret"],
      [[h[0], h[1], h[2]] for h in D["HIZMETLER"]], "4B2E83")
tablo(["Çalışma Saatleri", ""], [[s[0], s[1]] for s in D["SAATLER"]], "4B2E83")

serit("İLETİŞİM", "1F6F5C")
il = D["ILETISIM"]
tablo(["Kanal", "Bilgi"],
      [["Cep / WhatsApp", il["cep"]],
       ["İş Telefonu", il["isTel"]],
       ["E-posta", il["eposta"]],
       ["Adres", il["adres"]]])
for s in D["SOSYAL"]:
    satir_yaz("%s — %s" % (s["ad"], s["hesap"]), False, KOYU, 2, 10.5)

for s in doc.sections:
    f = s.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run("ÜSTAD KENAN KUZUCU · Resmî Biyografi · Sayfa ")
    r.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    f._p.append(fld)

doc.save(CIKTI)
print("yazıldı:", CIKTI, round(os.path.getsize(CIKTI) / 1024, 1), "KB")
